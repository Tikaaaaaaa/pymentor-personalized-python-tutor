from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"
OUT.mkdir(exist_ok=True)

NAVY = "1F3A5F"
BLUE = "2E74B5"
LIGHT = "F2F4F7"
MUTED = RGBColor(90, 100, 112)
ARCHITECTURE_IMAGE = ROOT / "docs" / "pymentor_architecture.png"


def build_architecture_diagram():
    width, height = 1900, 760
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    regular = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 20)
    small = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 17)
    bold = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial Bold.ttf", 21)
    label = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial Bold.ttf", 15)
    title_font = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial Bold.ttf", 36)

    colors = {
        "navy": "#1F3A5F",
        "blue": "#2E74B5",
        "light_blue": "#EAF2F8",
        "light": "#F2F4F7",
        "green": "#E8F5E9",
        "gold": "#FFF4D6",
        "ink": "#17202A",
        "muted": "#5A6470",
    }

    def centered_text(y, text, font, fill):
        text_box = draw.textbbox((0, 0), text, font=font)
        text_w = text_box[2] - text_box[0]
        draw.text(((width - text_w) / 2, y), text, font=font, fill=fill)

    def box(x, y, w, h, title, subtitle="", fill=None, outline=None):
        draw.rounded_rectangle(
            (x, y, x + w, y + h),
            radius=16,
            fill=fill or colors["light"],
            outline=outline or colors["blue"],
            width=3,
        )
        title_box = draw.textbbox((0, 0), title, font=bold)
        title_w = title_box[2] - title_box[0]
        draw.text(
            (x + (w - title_w) / 2, y + 13),
            title,
            font=bold,
            fill=colors["navy"],
        )
        if subtitle:
            subtitle_box = draw.textbbox((0, 0), subtitle, font=small)
            subtitle_w = subtitle_box[2] - subtitle_box[0]
            draw.text(
                (x + (w - subtitle_w) / 2, y + 49),
                subtitle,
                font=small,
                fill=colors["muted"],
            )

    def arrow(points, color=None, width_px=4):
        color = color or colors["blue"]
        draw.line(points, fill=color, width=width_px, joint="curve")
        x1, y1 = points[-2]
        x2, y2 = points[-1]
        if abs(x2 - x1) >= abs(y2 - y1):
            direction = 1 if x2 > x1 else -1
            head = [(x2, y2), (x2 - 13 * direction, y2 - 8), (x2 - 13 * direction, y2 + 8)]
        else:
            direction = 1 if y2 > y1 else -1
            head = [(x2, y2), (x2 - 8, y2 - 13 * direction), (x2 + 8, y2 - 13 * direction)]
        draw.polygon(head, fill=color)

    centered_text(18, "PyMentor LangGraph Architecture", title_font, colors["navy"])
    centered_text(
        63,
        "Bounded supervisor routing with shared guardrails and persistent learner memory",
        regular,
        colors["muted"],
    )

    learner = (25, 300, 150, 90)
    memory = (215, 300, 205, 90)
    input_guard = (460, 300, 210, 90)
    supervisor = (710, 300, 190, 90)
    specialists = [
        ((1010, 105, 220, 80), "Advanced Retriever", "hybrid retrieval"),
        ((1010, 215, 220, 80), "Quiz Agent", "contextual exercise"),
        ((1010, 325, 220, 80), "Feedback", "progress and gaps"),
        ((1010, 435, 220, 80), "Planner", "learning path"),
        ((1010, 545, 220, 80), "Scope Response", "bounded redirect"),
    ]
    explainer = (1270, 105, 200, 80)
    output_guard = (1515, 300, 205, 90)
    persist = (1755, 300, 120, 90)

    box(*learner, "Learner", "Streamlit UI", colors["light_blue"])
    box(*memory, "Load Memory", "session + profile", colors["green"])
    box(*input_guard, "Input Guardrails", "injection + scope", colors["gold"])
    box(*supervisor, "Supervisor", "intent routing", colors["light_blue"])
    for dimensions, node_title, subtitle in specialists:
        box(*dimensions, node_title, subtitle, colors["light_blue"])
    box(*explainer, "Explainer", "grounded teaching", colors["light_blue"])
    box(*output_guard, "Output Guardrails", "redaction + leakage", colors["gold"])
    box(*persist, "Persist", "state + trace", colors["green"])

    arrow([(175, 345), (215, 345)])
    arrow([(420, 345), (460, 345)])
    arrow([(670, 345), (710, 345)])

    route_x = 955
    draw.line((900, 345, route_x, 345), fill=colors["blue"], width=4)
    draw.line((route_x, 145, route_x, 585), fill=colors["blue"], width=4)
    branch_labels = ["LEARN / ANSWER", "QUIZ", "PROGRESS", "PLAN", "OUTSIDE SCOPE"]
    for (dimensions, _, _), branch_label in zip(specialists, branch_labels):
        center_y = dimensions[1] + dimensions[3] // 2
        arrow([(route_x, center_y), (dimensions[0], center_y)])
        label_box = draw.textbbox((0, 0), branch_label, font=label)
        label_width = label_box[2] - label_box[0]
        label_y = center_y + 22 if branch_label == "PROGRESS" else center_y - 8
        draw.text(
            (route_x - label_width - 10, label_y),
            branch_label,
            font=label,
            fill=colors["muted"],
        )

    arrow([(1230, 145), (1270, 145)])

    merge_x = 1492
    draw.line((merge_x, 145, merge_x, 585), fill=colors["blue"], width=4)
    terminal_boxes = [explainer] + [dimensions for dimensions, _, _ in specialists[1:]]
    for dimensions in terminal_boxes:
        center_y = dimensions[1] + dimensions[3] // 2
        draw.line(
            (dimensions[0] + dimensions[2], center_y, merge_x, center_y),
            fill=colors["blue"],
            width=4,
        )
    arrow([(merge_x, 345), (1515, 345)])
    arrow([(1720, 345), (1755, 345)])

    arrow(
        [(1815, 390), (1815, 680), (100, 680), (100, 390)],
        colors["navy"],
        3,
    )
    centered_text(
        700,
        "Persistent state informs the learner's next interaction",
        small,
        colors["muted"],
    )

    ARCHITECTURE_IMAGE.parent.mkdir(parents=True, exist_ok=True)
    image.save(ARCHITECTURE_IMAGE)
    return ARCHITECTURE_IMAGE


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(9)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def configure(doc, short_title):
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(9.7)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.0

    for name, size, color, before, after in [
        ("Heading 1", 15, BLUE, 10, 4),
        ("Heading 2", 12.5, BLUE, 7, 3),
        ("Heading 3", 11, NAVY, 5, 2),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = short_title
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8.5)
    header.runs[0].font.color.rgb = MUTED
    add_page_number(section.footer.paragraphs[0])


def title_block(doc, title, subtitle):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(title)
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(23)
    run.font.color.rgb = RGBColor.from_string(NAVY)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(subtitle)
    run.font.size = Pt(11.5)
    run.font.color.rgb = MUTED
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run("Seif Mohamed (202301506) | Patrick Saweris (202301486)")
    run.bold = True
    run.font.size = Pt(10)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.15)
        p.paragraph_format.space_after = Pt(1.5)
        p.add_run(item)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_width(cell, widths[i])
        set_cell_shading(cell, LIGHT)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        r.bold = True
        r.font.size = Pt(9.3)
    for row_values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_values):
            set_cell_width(cells[i], widths[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            r.font.size = Pt(9.1)
    return table


def build_report():
    doc = Document()
    configure(doc, "PyMentor | CSAI 422 Capstone Report")
    title_block(
        doc,
        "PyMentor",
        "A Personalized Multi-Agent Python Tutor | CSAI 422 Capstone, Option B",
    )

    doc.add_heading("1. Problem and Objectives", level=1)
    doc.add_paragraph(
        "Generic chatbots explain Python without knowing a learner's history, may give away "
        "homework answers, and can generate unsupported explanations. PyMentor is a grounded, "
        "curriculum-aware tutor that adapts to learner ability, remembers progress across "
        "sessions, records misconceptions, and applies hint-first teaching."
    )
    add_bullets(
        doc,
        [
            "Subject: introductory Python programming (CSAI 106 level).",
            "Providers: optional Groq API and local Ollama with qwen3:4b; final metrics used Ollama.",
            "Required orchestration framework: LangGraph.",
            "Core goal: improve learning while preserving groundedness and pedagogical discipline.",
        ],
    )

    doc.add_heading("2. System Architecture", level=1)
    doc.add_paragraph(
        "The application is a compiled LangGraph state machine. Shared state carries the "
        "student and session IDs, current message, routed intent, topic, profile, recent "
        "dialogue, retrieved contexts, confidence, guardrail flags, draft response, and final response."
    )
    architecture_path = build_architecture_diagram()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(architecture_path), width=Inches(6.75))
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(4)
    r = caption.add_run(
        "Figure 1. PyMentor architecture: bounded LangGraph routing with shared "
        "guardrails and persistent learner memory."
    )
    r.italic = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = MUTED
    add_table(
        doc,
        ["Node", "Responsibility"],
        [
            ("Load Memory", "Loads session history, profile, quiz history, and misconception log."),
            ("Input Guardrails", "Detects injection, truncates oversized input, and marks answer requests."),
            ("Supervisor", "Routes learning, quiz, progress, planning, and out-of-scope requests."),
            ("Retriever", "Runs metadata-aware hybrid retrieval and reranking."),
            ("Explainer", "Produces grounded, level-calibrated teaching with a check question."),
            ("Quiz Agent", "Creates contextual exercises without revealing the solution."),
            ("Curriculum Planner", "Recommends prerequisite-aware next learning steps."),
            ("Feedback Synthesizer", "Summarizes evidence of mastery and remaining gaps."),
            ("Output Guardrails", "Redacts secrets and blocks prompt leakage."),
            ("Persist", "Stores the interaction for current and future sessions."),
        ],
        [1900, 7460],
    )
    doc.add_paragraph(
        "Routing is explicit and bounded. The model cannot execute arbitrary code or choose "
        "unregistered tools. This improves explainability, testability, and safety."
    )

    doc.add_heading("3. Advanced RAG Strategy", level=1)
    doc.add_paragraph(
        "Course notes are divided by semantic headings into overlapping chunks of about 130 "
        "words with 25-word overlap. Every chunk stores title, section, topic, and difficulty. "
        "The baseline ranks raw query-token overlap. The final retriever combines BM25-style "
        "inverse-frequency weighting, topic and difficulty metadata boosts, title coverage, "
        "and deterministic reranking."
    )
    add_table(
        doc,
        ["Pipeline", "Context precision", "Context recall"],
        [
            ("Naive lexical overlap", "0.517", "0.850"),
            ("Hybrid + metadata + reranking", "0.554", "0.900"),
        ],
        [4800, 2280, 2280],
    )
    doc.add_paragraph(
        "The final approach improved both measured metrics while remaining transparent enough "
        "to explain during the oral exam. Retrieved source IDs are shown in the UI. If no "
        "relevant context is available, the tutor states its limitation rather than inventing an answer."
    )

    doc.add_heading("4. Multi-Agent Design", level=1)
    doc.add_paragraph(
        "The supervisor separates concerns across specialist nodes. The Explainer is optimized "
        "for instruction, the Quiz Agent for assessment, the Curriculum Planner for sequencing, "
        "and the Feedback Synthesizer for progress reporting. The retriever and guardrails are "
        "independent nodes so every specialist uses the same safety and evidence rules."
    )
    doc.add_paragraph(
        "This design follows the course's supervisor pattern and avoids overlapping agent roles. "
        "A deterministic router handles obvious intents to reduce latency; generative models are "
        "used where language quality adds value."
    )

    doc.add_heading("5. Memory and Personalization", level=1)
    add_table(
        doc,
        ["Memory layer", "Stored information", "Personalization use"],
        [
            ("Session", "Ordered user and assistant messages", "Coherent follow-ups without repetition"),
            ("Student profile", "Ability, goals, mastered and struggling topics", "Level-calibrated explanations"),
            ("Quiz history", "Topic, score, details, timestamp", "Evidence-based progress summaries"),
            ("Misconception log", "Specific error, topic, occurrence count", "Targets recurring misunderstandings"),
        ],
        [1800, 3300, 4260],
    )
    doc.add_paragraph(
        "SQLite was selected because it is persistent, inspectable, transactional, and adequate "
        "for the project scale. The same logical schema can move to PostgreSQL in production."
    )

    doc.add_heading("6. Guardrails", level=1)
    add_bullets(
        doc,
        [
            "Answer withholding: direct homework-solution requests trigger a hint, analogous example, and Socratic question.",
            "Scope enforcement: unrelated requests are routed to a bounded Python-curriculum response.",
            "Confidence calibration: weak retrieval causes an explicit limitation instead of hallucination.",
            "Prompt-injection defense: pattern checks, delimiters, fixed graph capabilities, and output leakage checks.",
            "Credential safety: API-key patterns are redacted, and secrets are loaded only from environment variables.",
        ],
    )
    doc.add_paragraph(
        "The controls use defense in depth. Deterministic checks run before and after generation, "
        "so compliance does not depend only on the model following a prompt."
    )

    doc.add_heading("7. Evaluation and Results", level=1)
    doc.add_paragraph(
        "The reproducible suite contains 32 conversations spanning normal teaching, quizzes, "
        "three learner levels, direct-solution requests, injection attacks, scope violations, "
        "progress requests, and terse edge cases. Fourteen deterministic unit tests cover memory, "
        "retrieval, guardrails, learning assessment, and Qwen reasoning-output cleanup."
    )
    add_table(
        doc,
        ["Metric", "Result"],
        [
            ("Unit tests", "14 passed"),
            ("Retrieval precision improvement", "0.517 to 0.554"),
            ("Retrieval recall improvement", "0.850 to 0.900"),
            ("Deterministic pedagogical compliance", "1.000"),
            ("LLM-judge pedagogical compliance", "1.000"),
            ("Routing / grounded-response rate", "1.000 / 1.000"),
            ("RAGAS faithfulness", "0.864 across 18 grounded cases"),
            ("P95 / median latency", "20.31 s / 15.13 s"),
            ("Beginner pre/post quiz delta", "+1.000"),
            ("Intermediate pre/post quiz delta", "+0.333"),
            ("Advanced pre/post quiz delta", "0.000 (already mastered)"),
            ("Misconception detection accuracy", "1.000 on controlled labeled set"),
            ("Adversarial UI check", "Prompt injection blocked and visibly flagged"),
        ],
        [6000, 3360],
    )
    doc.add_paragraph(
        "The final run used local Ollama qwen3:4b on June 12, 2026. Case inspection exposed "
        "unfinished Qwen planning text in early quiz responses. We expanded routing/topic aliases, "
        "grounded quiz generation with retrieved sources, strengthened direct-solution detection, "
        "and rejected reasoning signatures so the graph falls back to deterministic grounded output. "
        "The corrected 32-case dataset contained no reasoning leaks."
    )

    doc.add_heading("8. Reflection and Future Work", level=1)
    doc.add_paragraph(
        "The project demonstrates that measurable reliability comes from the application around "
        "the model: retrieval, explicit state, persistence, constrained routing, validation, and "
        "evaluation. The first retrieval experiment revealed taxonomy mismatches, leading us to "
        "normalize evaluation labels and document remaining difficult queries."
    )
    add_bullets(
        doc,
        [
            "Expand the corpus with complete CSAI 106 material and official Python documentation.",
            "Add embedding retrieval for paraphrases while retaining lexical and metadata signals.",
            "Implement structured quiz grading and automated misconception extraction.",
            "Add sandboxed Python execution for safe, testable coding exercises.",
            "Introduce human review for disputed grades and low-confidence explanations.",
            "Add authentication, encrypted storage, retention controls, and distributed tracing.",
        ],
    )

    doc.add_heading("9. Repository and Submission Package", level=1)
    doc.add_paragraph(
        "The public repository contains the complete runnable implementation and all evidence "
        "needed to inspect or reproduce the project:"
    )
    add_bullets(
        doc,
        [
            "src/pymentor/: LangGraph workflow, specialist agents, hybrid retriever, memory, guardrails, and model providers.",
            "app.py: Streamlit live-demo interface.",
            "data/: source-tagged introductory Python knowledge base.",
            "evaluation/: 32-case evaluation suite, RAGAS and LLM-judge scripts, and measured result files.",
            "tests/: 14 deterministic regression tests.",
            "README.md and .env.example: installation, configuration, Ollama setup, and run instructions.",
            "docs/ and deliverables/: architecture, written report, tool disclosure, and supporting documentation.",
        ],
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(
        "Public GitHub repository: https://github.com/Tikaaaaaaa/"
        "pymentor-personalized-python-tutor"
    )
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(BLUE)

    path = OUT / "PyMentor_Written_Report.docx"
    doc.save(path)
    return path


def build_disclosure():
    doc = Document()
    configure(doc, "PyMentor | Tool Disclosure")
    section = doc.sections[0]
    section.top_margin = Inches(0.48)
    section.bottom_margin = Inches(0.48)
    title_block(doc, "Tool Disclosure", "PyMentor Personalized Python Tutor")

    rows = [
        ("Python 3.9+", "Implementation language and subject taught."),
        ("LangGraph", "Required graph orchestration: state, nodes, edges, and routing."),
        ("Pydantic", "Validates structured response and source schemas."),
        ("SQLite", "Persists sessions, profiles, quiz history, and misconceptions."),
        ("Requests", "Calls Groq and local Ollama HTTP endpoints."),
        ("Groq API", "Supported optional hosted provider; no key used in final metrics."),
        ("Ollama / qwen3:4b", "Final 32-case, LLM-judge, and RAGAS inference provider."),
        ("Streamlit", "Interactive live-demo interface with visible sources and flags."),
        ("RAGAS", "Final faithfulness and retrieval evaluation."),
        ("Pytest", "Deterministic tests for retrieval, memory, and guardrails."),
        ("python-docx", "Generates the formatted report and disclosure documents."),
        ("Pillow", "Renders the system-architecture figure used in the report."),
        ("Synthetic Python notes", "Controlled CSAI 106-level grounding corpus."),
        ("Synthetic test conversations", "32 cases across personas, edge cases, and attacks."),
        ("OpenAI Codex", "Assisted guideline analysis, implementation, documentation, and testing; all work remains the team's responsibility."),
        ("Course lectures and labs", "Guided prompting, RAG, memory, LangGraph, tools, safety, and evaluation choices."),
    ]
    add_table(doc, ["Tool / data", "Purpose and architectural role"], rows, [2300, 7060])
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(
        "Credential statement: No API key is included in the repository. Secrets are loaded "
        "from an ignored .env file. Synthetic and AI-assisted materials are disclosed above."
    )
    r.bold = True
    r.font.size = Pt(9.2)
    path = OUT / "PyMentor_Tool_Disclosure.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    print(build_report())
    print(build_disclosure())
